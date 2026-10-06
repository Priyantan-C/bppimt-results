import json, os, re, sys, glob, shutil
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
JS, OUT = sys.argv[1], sys.argv[2]
shutil.rmtree(OUT, ignore_errors=True)
thin = Side(style="thin", color="BFBFBF"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
navy = PatternFill("solid", fgColor="1F3864")

def style(ws, widths=None):
    for col in ws.iter_cols():
        w = 0
        for c in col:
            c.border = border; c.number_format = "@"
            c.alignment = Alignment(horizontal="left" if c.column == 2 else "center", vertical="center")
            w = max(w, len(str(c.value or "")))
        h = col[0]; h.font = Font(bold=True, color="FFFFFF"); h.fill = navy
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(h.column)].width = min(max(w + 3, 9), 45)
    ws.row_dimensions[1].height = 30; ws.freeze_panes = "C2"

report = {r["slug"]: r for r in json.load(open(os.path.join(JS, "_report.json")))}
rows_summary, actions = [], []
for f in sorted(glob.glob(os.path.join(JS, "*-semester-results-*.json"))):
    d = json.load(open(f)); r = report[d["slug"]]
    notes = r["notes"]
    wrong = sorted({n for n in notes if n.startswith("WRONGHEAD")})
    review = sorted({n for n in notes if n.startswith("REVIEW")})
    status = "Needs original Excel" if wrong else ("Check flagged rows" if review else "Ready")
    folder = os.path.join(OUT, d["year"]); os.makedirs(folder, exist_ok=True)
    name = f"{d['year']}_SEM{d['semester']}_RESULT.xlsx"
    wb = Workbook(); wb.remove(wb.active); extra = []
    for c in d["courses"]:
        for i, t in enumerate(c["tables"]):
            title = c["course"] if i == 0 else f"{c['course']} ({i + 1})"
            ws = wb.create_sheet(title); ws.append(t["headings"])
            for row in t["rows"]: ws.append(row)
            style(ws)
        if len(c["tables"]) > 1:
            extra.append(f"{c['course']}: " + " + ".join(str(len(t['rows'])) for t in c["tables"]))
    wb.save(os.path.join(folder, name))
    rows_summary.append([d["year"], f"{d['semester']}", "PDF" if r.get("pdf") else "Pasted tables", r["src_tables"],
                         sum(len(c["tables"]) for c in d["courses"]), r["src_rows"], r["out_rows"],
                         "Yes" if r["src_rows"] == r["out_rows"] else "NO", status, "; ".join(extra),
                         f"https://bppimt.ac.in/result/{d['slug']}/", f"{d['year']}/{name}"])
    for n in wrong: actions.append([d["year"], d["semester"], "Wrong subject headings on the site", n.replace("WRONGHEAD ", "")])
    for n in review: actions.append([d["year"], d["semester"], "Row to check", n.replace("REVIEW ", "")])
    for n in sorted({n for n in notes if n.startswith("INFO roll")}): actions.append([d["year"], d["semester"], "Rebuilt row, please verify", n.replace("INFO ", "")])
rows_summary.sort(key=lambda x: (x[0], int(x[1])))
wb = Workbook(); ws = wb.active; ws.title = "Summary"
ws.append(["ACADEMIC YEAR", "SEM", "SOURCE", "TABLES BEFORE", "TABLES AFTER", "STUDENT ROWS IN SOURCE", "STUDENT ROWS IN OUTPUT",
           "ALL ROWS KEPT", "STATUS", "COURSES WITH A SEPARATE OLD-SUBJECTS TABLE (rows)", "LIVE PAGE", "CLEANED FILE"])
for r in rows_summary: ws.append([str(x) for x in r])
style(ws); ws.freeze_panes = "C2"
ws2 = wb.create_sheet("Needs action"); ws2.append(["ACADEMIC YEAR", "SEM", "ISSUE", "DETAIL"])
for a in actions: ws2.append([str(x) for x in a])
style(ws2); ws2.column_dimensions["D"].width = 110
for c in ws2["D"]: c.alignment = Alignment(horizontal="left", vertical="center")
wb.save(os.path.join(OUT, "BPPIMT_results_unification_report.xlsx"))
from collections import Counter
print(Counter(r[8] for r in rows_summary), "rows", sum(r[6] for r in rows_summary), "tables", sum(r[3] for r in rows_summary), "->", sum(r[4] for r in rows_summary))
for a in actions: print(a)
print("extra-table courses:", sum(1 for r in rows_summary if r[9]), "posts")
